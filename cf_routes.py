from flask import Blueprint, render_template, request, session, redirect, url_for, jsonify, make_response
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import io

cf_bp = Blueprint('cf_bp', __name__, template_folder='templates')

# నెలల పేర్ల మ్యాపింగ్
MONTH_MAP = {
    1: "June-2026", 2: "July-2026", 3: "August-2026", 4: "September-2026",
    5: "October-2026", 6: "November-2026", 7: "December-2026", 8: "January-2027",
    9: "February-2027", 10: "March-2027", 11: "April-2027", 12: "May-2027"
}

# 📋 CF వారీగా కేటాయించిన బ్యాంకుల మ్యాపింగ్ (CF Allotment Mapping)
CF_BANK_MAPPING = {
    "NARESH": ["BANK OF BARODA", "BANK OF INDIA"],
    "SATHISH": ["KDCCB", "PUNJAB NATIONAL BANK"],
    "NAVEEN": ["CANARA BANK", "CENTRAL BANK OF INDIA"],
    "SRIKANTH": ["12903_STATE BANK OF INDIA", "20140_STATE BANK OF INDIA", "20366_STATE BANK OF INDIA"],
    "DEVADAS": ["SIRCILLA TELANGANA GRAMEENA BANK", "THANGALLAPALLY TELANGANA GRAMEENA BANK"],
    "NAGARJUNA": ["UNION BANK OF INDIA"],
    "KANAKARAJU": ["INDIAN BANK"],
    "KAMAL": ["INDIAN BANK"],
    "MARKANDEYULU": ["BOINPARI UNION BANK", "NIMMAPALLY"]
}

# 1. CF లాగిన్ / నేమ్ సెలెక్షన్ పేజీ
@cf_bp.route('/cf/login', methods=['GET', 'POST'])
def cf_login():
    if request.method == 'POST':
        data = request.get_json() if request.is_json else request.form
        cf_name = str(data.get('cf_name', '')).strip().upper()
        
        if cf_name:
            session['cf_logged_in'] = True
            session['cf_name'] = cf_name
            return jsonify({"status": "success", "message": "లాగిన్ విజయవంతమైంది!"})
        else:
            return jsonify({"status": "error", "message": "దయచేసి CF పేరును ఎంచుకోండి!"})
            
    return render_template('cf_login.html')

# 2. CF డాష్‌బోర్డ్
@cf_bp.route('/cf/dashboard')
def cf_dashboard():
    if not session.get('cf_logged_in'):
        return redirect(url_for('cf_bp.cf_login'))
    
    cf_name = session.get('cf_name')
    return render_template('cf_dashboard.html', cf_name=cf_name, months=MONTH_MAP)

# 3. విధానం 1: పేరు / ఏరియా ప్రకారం రిపోర్ట్ డౌన్‌లోడ్ రౌట్
@cf_bp.route('/cf/download-report', methods=['POST'])
def cf_download_report_by_name():
    if not session.get('cf_logged_in'):
        return jsonify({"status": "error", "message": "దయచేసి లాగిన్ అవ్వండి!"})
        
    try:
        from app import get_sheets_client, SPREADSHEET_NAME
        
        cf_name = session.get('cf_name')
        current_month = int(request.form.get('month_number', 1))
            
        client = get_sheets_client()
        spreadsheet = client.open(SPREADSHEET_NAME)
        
        df_trans = pd.DataFrame(spreadsheet.worksheet('Transactions_Year1').get_all_records())
        df_thrift = pd.DataFrame(spreadsheet.worksheet('NEW THRIFT DATA').get_all_records())
        
        df_trans['Govt_Share'] = pd.to_numeric(df_trans['Govt_Share'], errors='coerce').fillna(0)
        df_trans['MONTH_NUMBER'] = pd.to_numeric(df_trans['MONTH_NUMBER'], errors='coerce').fillna(0)
        
        df_trans_month = df_trans[df_trans['MONTH_NUMBER'] == current_month].copy()
        df_trans_month = df_trans_month.drop_duplicates(subset=['MONTH_NUMBER', 'RD ACCOUNT_1 (Beneficiary Contribution)'], keep='first')
        
        df_thrift['RD 1_Clean'] = df_thrift['RD 1'].astype(str).str.strip()
        df_trans_month['RD1_Clean'] = df_trans_month['RD ACCOUNT_1 (Beneficiary Contribution)'].astype(str).str.strip()
        
        df_cross = pd.merge(df_thrift, df_trans_month[['RD1_Clean', 'Govt_Share', 'MONTH_NUMBER', 'FILE NAME']], left_on='RD 1_Clean', right_on='RD1_Clean', how='left')
        df_cross['Govt_Share'] = df_cross['Govt_Share'].fillna(0)
        
        cf_col = 'NAME OF CF' if 'NAME OF CF' in df_cross.columns else 'CF NAME'
        group_df = df_cross[df_cross[cf_col].astype(str).str.strip().str.upper() == cf_name].copy()
        
        if group_df.empty:
            return "మీ పరిధిలో ఎలాంటి డేటా లభించలేదు.", 404
            
        return generate_excel_response(group_df, cf_name, current_month, "Name_Based")

    except Exception as e:
        return f"రిపోర్ట్ జనరేట్ చేయడంలో లోపం: {str(e)}"


# 4. విధానం 2: కేటాయించిన బ్యాంకుల (Allotted Bank) ప్రకారం రిపోర్ట్ డౌన్‌లోడ్ రౌట్
@cf_bp.route('/cf/download-report-allotted-bank', methods=['POST'])
def cf_download_report_by_bank():
    if not session.get('cf_logged_in'):
        return jsonify({"status": "error", "message": "దయచేసి లాగిన్ అవ్వండి!"})
        
    try:
        from app import get_sheets_client, SPREADSHEET_NAME
        
        cf_name = session.get('cf_name')
        current_month = int(request.form.get('month_number', 1))
            
        client = get_sheets_client()
        spreadsheet = client.open(SPREADSHEET_NAME)
        
        df_trans = pd.DataFrame(spreadsheet.worksheet('Transactions_Year1').get_all_records())
        df_thrift = pd.DataFrame(spreadsheet.worksheet('NEW THRIFT DATA').get_all_records())
        
        df_trans['Govt_Share'] = pd.to_numeric(df_trans['Govt_Share'], errors='coerce').fillna(0)
        df_trans['MONTH_NUMBER'] = pd.to_numeric(df_trans['MONTH_NUMBER'], errors='coerce').fillna(0)
        
        df_trans_month = df_trans[df_trans['MONTH_NUMBER'] == current_month].copy()
        df_trans_month = df_trans_month.drop_duplicates(subset=['MONTH_NUMBER', 'RD ACCOUNT_1 (Beneficiary Contribution)'], keep='first')
        
        df_thrift['RD 1_Clean'] = df_thrift['RD 1'].astype(str).str.strip()
        df_trans_month['RD1_Clean'] = df_trans_month['RD ACCOUNT_1 (Beneficiary Contribution)'].astype(str).str.strip()
        
        df_cross = pd.merge(df_thrift, df_trans_month[['RD1_Clean', 'Govt_Share', 'MONTH_NUMBER']], left_on='RD 1_Clean', right_on='RD1_Clean', how='left')
        df_cross['Govt_Share'] = df_cross['Govt_Share'].fillna(0)
        
        allowed_banks = CF_BANK_MAPPING.get(cf_name, [])
        
        def match_bank(row):
            b_name = str(row.get('FILE NAME') or row.get('File Name') or '').strip().upper()
            for allowed in allowed_banks:
                if allowed in b_name or b_name in allowed:
                    return True
            return False

        group_df = df_cross[df_cross.apply(match_bank, axis=1)].copy()
        
        if group_df.empty:
            return "మీ పరిధిలోని బ్యాంకులకు సంబంధించిన ఎలాంటి డేటా లభించలేదు.", 404
            
        return generate_excel_response(group_df, cf_name, current_month, "Bank_Allotted")

    except Exception as e:
        return f"రిపోర్ట్ జనరేట్ చేయడంలో లోపం: {str(e)}"


# 💡 కామన్ ఎక్సెల్ జనరేషన్ హెల్పర్ ఫంక్షన్ (కోడ్ రిపీటేషన్ తగ్గించడానికి)
def generate_excel_response(group_df, cf_name, current_month, report_type):
    df_active = group_df[group_df['Govt_Share'] > 0].copy()
    df_missing = group_df[group_df['Govt_Share'] == 0].copy()
    
    wb = openpyxl.Workbook()
    
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    blue_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    red_fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    green_fill = PatternFill(start_color='008000', end_color='008000', fill_type='solid')
    
    # --- Sheet 1: Active ---
    ws1 = wb.active
    ws1.title = 'Sent_To_Bank_List'
    ws1.append(['BANK_SNO', 'Name Of Benificiary', 'Aadhar Number', 'Bank Account Number', 'RD 1','FILE NAME', 'Govt Share Amount', 'Status'])
    for col_num in range(1, 8):
        ws1.cell(row=1, column=col_num).font = header_font
        ws1.cell(row=1, column=col_num).fill = blue_fill
        
    for _, r in df_active.iterrows():
        ws1.append([
            r.get('BANK_SNO', ''), r.get('Name Of Benificiary', ''), r.get('Aadhar Number', ''),
            r.get('Bank Account Number', ''), r.get('RD 1', ''), r.get('FILE NAME',''), r.get('Govt_Share', 0),
            f"{MONTH_MAP[current_month]} Govt_Share Released"
        ])
        
    # --- Sheet 2: Missing ---
    ws2 = wb.create_sheet(title='Missing_or_Zero_List')
    ws2.append(['BANK_SNO', 'Name Of Benificiary', 'Aadhar Number', 'Bank Account Number', 'RD 1','FILE NAME', 'Status'])
    for col_num in range(1, 7):
        ws2.cell(row=1, column=col_num).font = header_font
        ws2.cell(row=1, column=col_num).fill = red_fill
        
    for _, r in df_missing.iterrows():
        ws2.append([
            r.get('BANK_SNO', ''), r.get('Name Of Benificiary', ''), r.get('Aadhar Number', ''),
            r.get('Bank Account Number', ''), r.get('RD 1', ''),r.get('FILE NAME',''),
            f"Zero or missing transactions for {MONTH_MAP[current_month]}"
        ])
        
    # --- Sheet 3: All ---
    ws3 = wb.create_sheet(title='All_Beneficiaries_Master')
    ws3.append(['BANK_SNO', 'Name Of Benificiary', 'Aadhar Number', 'Bank Account Number', 'RD 1','FILE NAME', 'Govt_Share'])
    for col_num in range(1, 7):
        ws3.cell(row=1, column=col_num).font = header_font
        ws3.cell(row=1, column=col_num).fill = green_fill
        
    for _, r in group_df.iterrows():
        ws3.append([
            r.get('BANK_SNO', ''), r.get('Name Of Benificiary', ''), r.get('Aadhar Number', ''),
            r.get('Bank Account Number', ''), r.get('RD 1', ''),r.get('FILE NAME',''), r.get('Govt_Share', 0)
        ])
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = make_response(output.read())
    response.headers["Content-Disposition"] = f"attachment; filename={cf_name}_{report_type}_Month_{current_month}_Report.xlsx"
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response

@cf_bp.route('/cf/logout', methods=['POST'])
def cf_logout():
    session.pop('cf_logged_in', None)
    session.pop('cf_name', None)
    return jsonify({"status": "success", "message": "లాగౌట్ అయ్యారు."})
